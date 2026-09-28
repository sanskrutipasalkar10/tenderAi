"""app.pipeline.fetch_links — link extraction/classification (pure), fetching (HTTP
mocked, CLAUDE.md hard rule 8), and content-type-specific extraction. Uses a tiny
synthetic PDF built with fitz for the link-extraction test rather than any real
GeM-Bidding-*.pdf sample — those are real, gitignored tender documents
(docs/DECISIONS.md #27), never something a committed test can depend on existing.
"""

import io
import uuid
from unittest.mock import MagicMock, patch

import fitz
import pytest
import requests
from docx import Document as DocxDocument

from app.pipeline import fetch_links


def _pdf_with_link(url: str) -> fitz.Document:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "See Annexure A for the real scope of work.")
    page.insert_link({"kind": fitz.LINK_URI, "uri": url, "from": fitz.Rect(72, 72, 300, 90)})
    return doc


# --- extract_uri_links ---------------------------------------------------------------


def test_extract_uri_links_finds_a_real_uri_link() -> None:
    doc = _pdf_with_link("https://bidplus.gem.gov.in/bidding/bid/downloadMseMiiDoc/1/x.pdf")

    refs = fetch_links.extract_uri_links(doc)

    assert len(refs) == 1
    assert refs[0].url == "https://bidplus.gem.gov.in/bidding/bid/downloadMseMiiDoc/1/x.pdf"
    assert refs[0].page_number == 0
    assert "Annexure" in refs[0].text_near_link


def test_extract_uri_links_dedupes_the_same_url_across_pages() -> None:
    doc = fitz.open()
    url = "https://admin.gem.gov.in/apis/v1/gtc/pdfByDate/?date=20260101"
    for _ in range(3):
        page = doc.new_page()
        page.insert_link({"kind": fitz.LINK_URI, "uri": url, "from": fitz.Rect(72, 72, 300, 90)})

    refs = fetch_links.extract_uri_links(doc)

    assert len(refs) == 1


def test_extract_uri_links_returns_empty_for_a_pdf_with_no_links() -> None:
    doc = fitz.open()
    doc.new_page()

    assert fetch_links.extract_uri_links(doc) == []


# --- classify_link, real URL patterns confirmed against 3 real GeM-Bidding-*.pdf files -


@pytest.mark.parametrize(
    "url",
    [
        "https://bidplus.gem.gov.in/resources/upload_nas/JulQ325/bidding/biddoc/bid-8051120/1751693557.pdf",
        "https://bidplus.gem.gov.in/bidding/bid/downloadMseMiiDoc/8051120/1751694222.pdf",
        "https://bidplus.gem.gov.in/bidding/bid/bidsla/64820533254400",
        "https://fulfilment.gem.gov.in/contract/slafds?fileDownloadPath=x",
    ],
)
def test_classify_link_bid_specific(url: str) -> None:
    assert fetch_links.classify_link(url) == "bid_specific"


@pytest.mark.parametrize(
    "url",
    [
        "https://admin.gem.gov.in/apis/v1/gtc/pdfByDate/?date=20250705",
        "https://assets-bg.gem.gov.in/resources/upload/shared_doc/list-of-categories-where-trials-are-allowed_1712126171.pdf",
        "https://bidplus.gem.gov.in/bidding/downloadOmppdfile/",
    ],
)
def test_classify_link_boilerplate(url: str) -> None:
    assert fetch_links.classify_link(url) == "boilerplate"


def test_classify_link_unrecognized_pattern_is_unknown() -> None:
    assert fetch_links.classify_link("https://gem.gov.in/some/new/path/never/seen") == "unknown"


def test_is_boilerplate_matches_classify_link() -> None:
    assert fetch_links.is_boilerplate("https://bidplus.gem.gov.in/bidding/downloadOmppdfile/")
    assert not fetch_links.is_boilerplate(
        "https://bidplus.gem.gov.in/bidding/bid/bidsla/64820533254400"
    )


# --- fetch_url — HTTP mocked ----------------------------------------------------------


@patch("app.pipeline.fetch_links.is_fetchable_url", return_value=True)
@patch("app.pipeline.fetch_links.requests.get")
def test_fetch_url_returns_content_and_sniffed_type_for_pdf(mock_get, _mock_gate) -> None:
    mock_get.return_value = MagicMock(status_code=200, content=b"%PDF-1.4 fake pdf bytes")
    mock_get.return_value.raise_for_status = MagicMock()

    result = fetch_links.fetch_url("https://bidplus.gem.gov.in/x.pdf")

    assert result == (b"%PDF-1.4 fake pdf bytes", "pdf")


@patch("app.pipeline.fetch_links.is_fetchable_url", return_value=True)
@patch("app.pipeline.fetch_links.requests.get")
def test_fetch_url_sniffs_html_even_when_url_looks_like_a_document(mock_get, _mock_gate) -> None:
    # The real GeM "Service Level Agreement" link returns a full HTML portal page, not
    # a document — confirmed by direct inspection, this is the case that motivated
    # sniffing content by magic bytes rather than trusting the URL's apparent purpose.
    mock_get.return_value = MagicMock(
        status_code=200, content=b"<!DOCTYPE html><html><body>SLA text</body></html>"
    )
    mock_get.return_value.raise_for_status = MagicMock()

    result = fetch_links.fetch_url("https://bidplus.gem.gov.in/bidding/bid/bidsla/1")

    assert result is not None
    assert result[1] == "html"


@patch("app.pipeline.fetch_links.is_fetchable_url", return_value=False)
def test_fetch_url_returns_none_when_guardrail_rejects(_mock_gate) -> None:
    assert fetch_links.fetch_url("https://evil.example.com/x.pdf") is None


@patch("app.pipeline.fetch_links.is_fetchable_url", return_value=True)
@patch("app.pipeline.fetch_links.requests.get", side_effect=requests.ConnectionError("down"))
@patch("app.pipeline.fetch_links.time.sleep")  # no real sleeping in a test
def test_fetch_url_returns_none_after_retries_exhausted(_sleep, mock_get, _mock_gate) -> None:
    assert fetch_links.fetch_url("https://bidplus.gem.gov.in/x.pdf", num_retries=1) is None
    assert mock_get.call_count == 2  # initial attempt + 1 retry


@patch("app.pipeline.fetch_links.is_fetchable_url", return_value=True)
@patch("app.pipeline.fetch_links.requests.get")
def test_fetch_url_returns_none_for_unrecognized_content(mock_get, _mock_gate) -> None:
    mock_get.return_value = MagicMock(status_code=200, content=b"random unrecognized bytes")
    mock_get.return_value.raise_for_status = MagicMock()

    assert fetch_links.fetch_url("https://bidplus.gem.gov.in/x") is None


# --- HTML / docx extraction ------------------------------------------------------------


def test_extract_html_content_strips_nav_and_script_chrome() -> None:
    html = b"""
    <html><head><script>var x = 1;</script></head>
    <body>
      <nav>Login Sign Up Buyers Sellers</nav>
      <div id="content">Service Level Agreement clause: penalty 1% per week.</div>
      <footer>Copyright GeM</footer>
    </body></html>
    """
    text = fetch_links._extract_html_content(html)

    assert text is not None
    assert "penalty 1% per week" in text
    assert "Login Sign Up" not in text


def test_extract_docx_content_reads_paragraphs_and_tables() -> None:
    doc = DocxDocument()
    doc.add_paragraph("Clause 1: EMD is Rs 2,80,000.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Delivery period"
    table.rows[0].cells[1].text = "30 days"
    buf = io.BytesIO()
    doc.save(buf)

    text = fetch_links._extract_docx_content(buf.getvalue())

    assert text is not None
    assert "EMD is Rs 2,80,000" in text
    assert "Delivery period" in text and "30 days" in text


def test_text_to_page_record_empty_text_is_zero_confidence() -> None:
    record = fetch_links.text_to_page_record(None)
    assert record.raw_text is None
    assert record.confidence_score == 0.0


def test_text_to_page_record_real_text_is_full_confidence() -> None:
    record = fetch_links.text_to_page_record("real extracted text")
    assert record.raw_text == "real extracted text"
    assert record.confidence_score == 1.0
    assert record.content_hash is not None


# --- URL-keyed cache — mocked DB -------------------------------------------------------


class _FakeCacheSession:
    def __init__(self, existing=None) -> None:
        self._existing = existing
        self.added: list = []
        self.commit_count = 0

    def get(self, model, pk):
        return self._existing if self._existing and self._existing.url_hash == pk else None

    def add(self, obj) -> None:
        self.added.append(obj)

    def commit(self) -> None:
        self.commit_count += 1


def test_get_cached_pages_returns_none_on_miss() -> None:
    db = _FakeCacheSession(existing=None)
    assert fetch_links.get_cached_pages(db, "https://example.gem.gov.in/x.pdf") is None


def test_get_cached_pages_returns_pages_and_increments_hit_count_on_hit() -> None:
    from app.models.linked_document_cache import LinkedDocumentCache

    url = "https://admin.gem.gov.in/apis/v1/gtc/pdfByDate/?date=20260101"
    cached_record = fetch_links.text_to_page_record("General Terms and Conditions text")
    row = LinkedDocumentCache(
        url_hash=fetch_links._url_hash(url),
        source_url=url,
        cached_extraction=[cached_record.model_dump(mode="json")],
        hit_count=0,
    )
    db = _FakeCacheSession(existing=row)

    pages = fetch_links.get_cached_pages(db, url)

    assert pages is not None
    assert pages[0].raw_text == "General Terms and Conditions text"
    assert row.hit_count == 1
    assert db.commit_count == 1


def test_store_cached_pages_inserts_a_new_row() -> None:
    db = _FakeCacheSession()
    pages = [fetch_links.text_to_page_record("some boilerplate text")]

    fetch_links.store_cached_pages(db, "https://admin.gem.gov.in/gtc/x", uuid.uuid4(), pages)

    assert len(db.added) == 1
    assert db.added[0].source_url == "https://admin.gem.gov.in/gtc/x"
    assert db.added[0].cached_extraction[0]["raw_text"] == "some boilerplate text"
